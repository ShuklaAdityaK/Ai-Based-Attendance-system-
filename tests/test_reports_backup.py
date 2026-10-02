from datetime import date, datetime, timedelta
from io import BytesIO

from openpyxl import load_workbook

from core import attendance as A
from core import backup, db
from reports import excel_report, pdf_report
from reports.common import ReportFilters


def _history(env, n=6):
    admin, students = env["admin"], env["students"]
    sid = db.create_subject(admin, "CS704", "Cloud Computing")
    db.set_subject_enrollment(admin, sid, [s.id for s in students])
    for i in range(n):
        d = date(2026, 8, 3) + timedelta(days=i)
        sess = A.create_session(admin, sid, d, "09:00")
        A.open_session(admin, sess)
        start = datetime.combine(d, datetime.min.time()).replace(hour=9)
        A.mark_from_kiosk(admin, sess, students[0].id, 0.3, 0.95, 0, check_in=start + timedelta(minutes=2))
        if i % 2:
            A.mark_from_kiosk(admin, sess, students[1].id, 0.3, 0.95, 0, check_in=start + timedelta(minutes=15))
        A.close_session(admin, sess)
    return sid


def test_excel_report_sheets(env):
    _history(env)
    wb = load_workbook(BytesIO(excel_report.generate(env["admin"], ReportFilters())))
    assert {"Session records", "Student summary", "Subject summary"} <= set(wb.sheetnames)


def test_pdf_report_generates(env):
    _history(env)
    pdf = pdf_report.generate(env["admin"], ReportFilters())
    assert pdf[:4] == b"%PDF" and len(pdf) > 2000


def test_student_report_contains_only_self(env):
    _history(env)
    s1 = env["students"][0]
    wb = load_workbook(BytesIO(excel_report.generate(s1, ReportFilters())))
    ws = wb["Session records"]
    names = {c.value for c in ws["E"] if c.value and c.value.startswith("Student")}
    assert names == {"Student 1"}


def test_backup_roundtrip(env):
    _history(env, n=2)
    admin = env["admin"]
    zip_path = backup.create_backup(admin)
    assert backup.validate_backup(zip_path)["app"] == backup.APP_ID
    db.create_subject(admin, "CS799", "Added after backup")
    backup.restore_backup(admin, zip_path)
    codes = {s["code"] for s in db.list_subjects(admin)}
    assert "CS799" not in codes and "CS704" in codes
